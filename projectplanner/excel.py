"""Export a project's activities to an .xlsx workbook."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .store import ProjectStore, format_dependency_ids

# (header, column width)
COLUMNS = [
    ("ID", 6),
    ("Group", 20),
    ("Title", 36),
    ("Depends On", 14),
    ("Required By", 14),
    ("Most Likely (weeks)", 18),
    ("Min Ratio", 12),
    ("Max Ratio", 12),
    ("Status", 13),
    ("Owner", 16),
    ("Scope", 80),
    ("Risks", 50),
    ("Created", 20),
    ("Updated", 20),
]


def export_activities(store: ProjectStore, path: str | Path) -> int:
    """Write every activity to `path`; return the number of activities written."""
    activities = store.list_activities()
    required_by: dict[int, list[int]] = {a.id: [] for a in activities}
    for a in activities:
        for dep in a.depends_on:
            required_by.setdefault(dep, []).append(a.id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Activities"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    top_wrap = Alignment(vertical="top", wrap_text=True)
    for row, a in enumerate(activities, start=2):
        values = [
            a.id,
            a.group_name,
            a.title,
            format_dependency_ids(a.depends_on),
            format_dependency_ids(required_by.get(a.id, [])),
            a.duration_weeks,
            a.min_duration_ratio,
            a.max_duration_ratio,
            a.status,
            a.owner,
            a.description_text.strip(),
            a.risks_text.strip(),
            a.created_at.replace("T", " "),
            a.updated_at.replace("T", " "),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = top_wrap
        ws.cell(row=row, column=6).number_format = "0"
        ws.cell(row=row, column=7).number_format = "0.00"
        ws.cell(row=row, column=8).number_format = "0.00"

    last = len(activities) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"

    total_row = last + 2
    ws.cell(row=total_row, column=5, value="Nominal project total").font = Font(bold=True)
    total = ws.cell(row=total_row, column=6, value=f"=SUM(F2:F{max(last, 2)})")
    total.font = Font(bold=True)
    total.number_format = "0"

    wb.save(path)
    return len(activities)
