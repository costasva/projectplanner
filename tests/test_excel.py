from openpyxl import load_workbook

from projectplanner.excel import export_activities
from projectplanner.store import Activity, ProjectStore


def test_export(tmp_path):
    store = ProjectStore()
    a = store.create(Activity(title="Design", duration_weeks=8, description_text="Draw it"))
    store.create(Activity(title="Build", duration_weeks=20, depends_on=[a.id]))
    path = tmp_path / "out.xlsx"
    assert export_activities(store, path) == 2

    ws = load_workbook(path).active
    rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
    assert rows[0][:9] == ("ID", "Group", "Title", "Depends On", "Required By", "Most Likely (weeks)", "Min Ratio", "Max Ratio", "Status")
    assert rows[1][:9] == (1, None, "Design", None, "2", 8, 0.5, 2, "Not started")
    assert rows[2][:9] == (2, None, "Build", "1", None, 20, 0.5, 2, "Not started")
    assert rows[1][10] == "Draw it"
    assert ws["F5"].value == "=SUM(F2:F3)"
